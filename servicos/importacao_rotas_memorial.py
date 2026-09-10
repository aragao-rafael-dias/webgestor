from __future__ import annotations

import json
import math
import re
import tempfile
import unicodedata
from dataclasses import asdict, dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

import requests
from docx import Document

from servicos.memoriais_docx import extrair_metadados_docx, texto_limpo


MAX_PONTOS_POR_TRECHO = 500
MAX_VERTICES_GEOMETRIA = 100_000


@dataclass
class PontoItinerario:
    ordem: int
    referencia: str
    lat: float
    lon: float
    lado_via: str = ""
    logradouro: str = ""
    direcao_seguir: str = ""

    @property
    def coordenadas(self) -> str:
        return f"({self.lat:.6f}, {self.lon:.6f})"

    def para_dict(self) -> dict[str, Any]:
        dados = asdict(self)
        dados["coordenadas"] = self.coordenadas
        return dados


@dataclass
class TrechoImportado:
    trecho: str
    pontos: list[PontoItinerario] = field(default_factory=list)
    geometria: dict[str, Any] | None = None
    distancia_roteada_km: float | None = None
    duracao_roteada_min: float | None = None
    distancias_entre_pontos_km: list[float] = field(default_factory=list)

    def para_dict(self) -> dict[str, Any]:
        return {
            "trecho": self.trecho,
            "pontos": [ponto.para_dict() for ponto in self.pontos],
            "geometria": self.geometria,
            "distancia_roteada_km": self.distancia_roteada_km,
            "duracao_roteada_min": self.duracao_roteada_min,
            "distancias_entre_pontos_km": list(self.distancias_entre_pontos_km),
        }


@dataclass
class MemorialRotaImportado:
    numero_linha: str
    nome_rota: str
    regiao: str
    linha: str
    km_ida_documento: float | None
    km_volta_documento: float | None
    trechos: dict[str, TrechoImportado]
    avisos: list[str] = field(default_factory=list)
    alertas_geometria: list[str] = field(default_factory=list)

    def para_dict(self) -> dict[str, Any]:
        return {
            "numero_linha": self.numero_linha,
            "nome_rota": self.nome_rota,
            "regiao": self.regiao,
            "linha": self.linha,
            "km_ida_documento": self.km_ida_documento,
            "km_volta_documento": self.km_volta_documento,
            "trechos": {
                nome: trecho.para_dict()
                for nome, trecho in self.trechos.items()
            },
            "avisos": list(self.avisos),
            "alertas_geometria": list(self.alertas_geometria),
            "requer_confirmacao": bool(self.alertas_geometria),
        }


def _sem_acentos(valor: str) -> str:
    normalizado = unicodedata.normalize("NFD", texto_limpo(valor))
    return "".join(
        caractere
        for caractere in normalizado
        if unicodedata.category(caractere) != "Mn"
    )


def nome_rota_banco(numero_linha: str) -> str:
    codigo = texto_limpo(numero_linha)
    codigo = re.sub(r"^Rota", "", codigo, flags=re.IGNORECASE)
    codigo = codigo.replace(".", "_")
    codigo = re.sub(r"^([A-Za-z]+)-(\d+)$", r"\1_\2", codigo)
    codigo = re.sub(r"\s+", "", codigo)
    if not codigo:
        raise ValueError("Número da linha não identificado no memorial.")
    if not re.fullmatch(r"[A-Za-z0-9_-]+", codigo):
        raise ValueError("O número da linha possui caracteres não reconhecidos.")
    return f"Rota{codigo}"


def normalizar_regiao_banco(valor: str) -> str:
    texto = texto_limpo(valor)
    if not texto:
        return ""

    chave = re.sub(r"\s+", " ", _sem_acentos(texto).upper()).strip()
    numero = re.fullmatch(r"(?:REGIAO\s*)?0*(\d+)", chave)
    if numero:
        return f"REGIÃO {int(numero.group(1))}"
    if chave in {"UNIVERSIDADE", "REGIAO UNIVERSIDADE"}:
        return "REGIÃO UNIVERSIDADE"
    if chave.startswith("REGIAO "):
        return "REGIÃO " + texto_limpo(chave[len("REGIAO "):])
    return texto


def _texto_itinerario(documento: Document) -> str:
    candidatos: list[str] = []
    for tabela in documento.tables:
        for linha in tabela.rows:
            vistos: set[str] = set()
            for celula in linha.cells:
                texto = celula.text or ""
                chave = texto.strip()
                if not chave or chave in vistos:
                    continue
                vistos.add(chave)
                chave_sem_acentos = _sem_acentos(chave).upper()
                if "ITINERARIO DA LINHA" in chave_sem_acentos:
                    candidatos.append(chave)
    if not candidatos:
        raise ValueError("O memorial não possui a seção 'Itinerário da Linha'.")
    return max(candidatos, key=len)


def _extrair_regiao(documento: Document) -> str:
    for tabela in documento.tables:
        for linha in tabela.rows[:2]:
            for celula in linha.cells:
                match = re.search(r"(?mi)^\s*Região\s*:\s*([^\n\r]+)", celula.text or "")
                if match:
                    return normalizar_regiao_banco(match.group(1))
    return ""


def _limpar_linha_itinerario(linha: str) -> str:
    linha = texto_limpo(linha)
    linha = re.sub(r"^[•·▪◦\-*]+\s*", "", linha)
    linha = re.sub(r"^[➡➜➔➤→]\ufe0f?\s*", "", linha)
    return linha.strip()


def _novo_ponto(match: re.Match[str]) -> PontoItinerario:
    return PontoItinerario(
        ordem=int(match.group(1)),
        referencia=texto_limpo(match.group(2)),
        lat=0.0,
        lon=0.0,
    )


def extrair_itinerarios(documento: Document) -> dict[str, TrechoImportado]:
    texto = _texto_itinerario(documento)
    trechos = {
        "IDA": TrechoImportado("IDA"),
        "VOLTA": TrechoImportado("VOLTA"),
    }
    trecho_atual: TrechoImportado | None = None
    ponto_atual: PontoItinerario | None = None

    def finalizar_ponto() -> None:
        nonlocal ponto_atual
        if ponto_atual is None or trecho_atual is None:
            return
        if not (-90 <= ponto_atual.lat <= 90 and -180 <= ponto_atual.lon <= 180):
            raise ValueError(
                f"Coordenada inválida no ponto {ponto_atual.ordem} de {trecho_atual.trecho}."
            )
        if ponto_atual.lat == 0.0 and ponto_atual.lon == 0.0:
            raise ValueError(
                f"Coordenada não identificada no ponto {ponto_atual.ordem} de {trecho_atual.trecho}."
            )
        trecho_atual.pontos.append(ponto_atual)
        ponto_atual = None

    for linha_original in texto.splitlines():
        linha = _limpar_linha_itinerario(linha_original)
        if not linha:
            continue

        cabecalho = re.sub(r"[^A-Z]", "", _sem_acentos(linha).upper())
        if cabecalho == "IDA":
            finalizar_ponto()
            trecho_atual = trechos["IDA"]
            continue
        if cabecalho == "VOLTA":
            finalizar_ponto()
            trecho_atual = trechos["VOLTA"]
            continue
        if trecho_atual is None:
            continue
        if cabecalho in {"INICIO", "TERMINO", "CHEGADA", "FIM"}:
            continue

        match_ponto = re.match(r"^(\d+)\s*[.)-]\s*(.+)$", linha)
        if match_ponto:
            finalizar_ponto()
            ponto_atual = _novo_ponto(match_ponto)
            continue
        if ponto_atual is None:
            continue

        match_coord = re.search(
            r"COORDENADAS?\s*:\s*\(\s*([-+]?\d+(?:[.,]\d+)?)\s*,\s*([-+]?\d+(?:[.,]\d+)?)\s*\)",
            linha,
            flags=re.IGNORECASE,
        )
        if match_coord:
            ponto_atual.lat = float(match_coord.group(1).replace(",", "."))
            ponto_atual.lon = float(match_coord.group(2).replace(",", "."))
            continue

        match_lado = re.match(r"^LADO\s*:\s*(.*)$", linha, flags=re.IGNORECASE)
        if match_lado:
            ponto_atual.lado_via = texto_limpo(match_lado.group(1)).upper()
            continue

        match_logradouro = re.match(r"^LOGRADOURO\s*:\s*(.*)$", linha, flags=re.IGNORECASE)
        if match_logradouro:
            ponto_atual.logradouro = texto_limpo(match_logradouro.group(1))
            continue

        match_direcao = re.match(
            r"^(?:DIREÇÃO(?: A SEGUIR)?|DIRECAO(?: A SEGUIR)?)\s*:\s*(.*)$",
            linha,
            flags=re.IGNORECASE,
        )
        if match_direcao:
            ponto_atual.direcao_seguir = texto_limpo(match_direcao.group(1))
            continue

        linha_direcao = re.sub(r"^[➡➜➔➤→]\ufe0f?\s*", "", linha).strip()
        if re.search(r"\bKM\b", linha_direcao, flags=re.IGNORECASE) or re.match(
            r"^(EM FRENTE|VIRANDO|SEGUINDO|RETORNANDO|CONTINUANDO)",
            linha_direcao,
            flags=re.IGNORECASE,
        ):
            ponto_atual.direcao_seguir = linha_direcao

    finalizar_ponto()

    for nome, trecho in trechos.items():
        if len(trecho.pontos) > MAX_PONTOS_POR_TRECHO:
            raise ValueError(f"O trecho {nome} excede o limite de {MAX_PONTOS_POR_TRECHO} pontos.")
        ordens = [p.ordem for p in trecho.pontos]
        if len(ordens) != len(set(ordens)):
            raise ValueError(f"Há números de ponto repetidos no trecho {nome}.")

    if not any(len(trecho.pontos) >= 2 for trecho in trechos.values()):
        raise ValueError(
            "O memorial precisa ter pelo menos um trecho (IDA ou VOLTA) com dois pontos coordenados."
        )

    return trechos


def _chave_referencia(valor: str) -> str:
    chave = _sem_acentos(valor).upper()
    chave = re.sub(r"[^A-Z0-9]+", " ", chave)
    return re.sub(r"\s+", " ", chave).strip()


def _referencias_equivalentes(a: str, b: str) -> bool:
    chave_a = _chave_referencia(a)
    chave_b = _chave_referencia(b)
    if not chave_a or not chave_b:
        return False
    if chave_a == chave_b:
        return True
    if chave_a in chave_b or chave_b in chave_a:
        return min(len(chave_a), len(chave_b)) / max(len(chave_a), len(chave_b)) >= 0.75
    return SequenceMatcher(None, chave_a, chave_b).ratio() >= 0.88


def distancia_haversine_km(a: PontoItinerario, b: PontoItinerario) -> float:
    raio = 6371.0088
    lat1 = math.radians(a.lat)
    lat2 = math.radians(b.lat)
    dlat = lat2 - lat1
    dlon = math.radians(b.lon - a.lon)
    h = math.sin(dlat / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlon / 2) ** 2
    return 2 * raio * math.asin(min(1.0, math.sqrt(h)))


def reconciliar_extremos(trechos: dict[str, TrechoImportado]) -> list[str]:
    ida = trechos["IDA"].pontos
    volta = trechos["VOLTA"].pontos
    avisos: list[str] = []

    if not ida or not volta:
        return avisos

    pares = (
        (ida[-1], volta[0], "término da IDA / início da VOLTA"),
        (ida[0], volta[-1], "início da IDA / término da VOLTA"),
    )
    for referencia, candidato, descricao in pares:
        distancia = distancia_haversine_km(referencia, candidato)
        if _referencias_equivalentes(referencia.referencia, candidato.referencia):
            if distancia > 0.25:
                lat_antiga, lon_antiga = candidato.lat, candidato.lon
                candidato.lat = referencia.lat
                candidato.lon = referencia.lon
                avisos.append(
                    f"{descricao}: a mesma referência aparecia com coordenadas separadas por "
                    f"{distancia:.1f} km. O ponto ({lat_antiga:.6f}, {lon_antiga:.6f}) foi "
                    f"alinhado para ({candidato.lat:.6f}, {candidato.lon:.6f}) antes do roteamento."
                )
        elif distancia > 5.0:
            avisos.append(
                f"{descricao}: os extremos estão separados por {distancia:.1f} km e têm referências "
                "diferentes. Confira a prévia antes de salvar."
            )
    return avisos


def _distancia_declarada_km(direcao: str) -> float | None:
    match = re.search(r"(\d+(?:[.,]\d+)?)\s*KM\b", direcao or "", flags=re.IGNORECASE)
    if not match:
        return None
    try:
        return float(match.group(1).replace(",", "."))
    except ValueError:
        return None


def roteamento_osrm(
    pontos: list[PontoItinerario],
    base_url: str = "https://router.project-osrm.org",
    timeout: int = 45,
) -> tuple[dict[str, Any], float, float, list[float]]:
    coordenadas = ";".join(f"{p.lon:.6f},{p.lat:.6f}" for p in pontos)
    url = f"{base_url.rstrip('/')}/route/v1/driving/{coordenadas}"
    resposta = requests.get(
        url,
        params={
            "overview": "full",
            "geometries": "geojson",
            "steps": "false",
        },
        headers={"User-Agent": "WebSIG-SEMED/1.0"},
        timeout=timeout,
    )
    resposta.raise_for_status()
    dados = resposta.json()
    if dados.get("code") != "Ok" or not dados.get("routes"):
        mensagem = dados.get("message") or dados.get("code") or "sem rota"
        raise ValueError(f"O roteador não conseguiu calcular o trajeto: {mensagem}.")

    rota = dados["routes"][0]
    geometria = rota.get("geometry")
    if not geometria or geometria.get("type") != "LineString":
        raise ValueError("O roteador retornou uma geometria inválida.")

    pernas = [
        round(float(perna.get("distance", 0.0)) / 1000.0, 3)
        for perna in (rota.get("legs") or [])
    ]
    if pernas and len(pernas) != len(pontos) - 1:
        raise ValueError("O roteador retornou uma quantidade inesperada de segmentos.")

    return (
        geometria,
        float(rota.get("distance", 0.0)) / 1000.0,
        float(rota.get("duration", 0.0)) / 60.0,
        pernas,
    )


def _avisos_distancias_pontos(nome: str, trecho: TrechoImportado) -> list[str]:
    avisos: list[str] = []
    for indice, distancia_roteada in enumerate(trecho.distancias_entre_pontos_km):
        ponto = trecho.pontos[indice]
        declarada = _distancia_declarada_km(ponto.direcao_seguir)
        if declarada is None or declarada <= 0:
            continue
        diferenca_abs = abs(distancia_roteada - declarada)
        diferenca_rel = diferenca_abs / declarada
        if diferenca_abs >= 1.0 and diferenca_rel >= 0.50:
            avisos.append(
                f"{nome}, ponto {ponto.ordem}: o memorial informa {declarada:.3f} km até o "
                f"próximo ponto, enquanto o roteador calculou {distancia_roteada:.3f} km."
            )
    return avisos


def importar_docx(
    caminho_docx: Path,
    osrm_base_url: str = "https://router.project-osrm.org",
    timeout: int = 45,
) -> MemorialRotaImportado:
    metadados = extrair_metadados_docx(caminho_docx)
    documento = Document(caminho_docx)
    trechos = extrair_itinerarios(documento)

    avisos = list(metadados.avisos)
    alertas_geometria = reconciliar_extremos(trechos)

    for nome, trecho in trechos.items():
        if len(trecho.pontos) < 2:
            avisos.append(
                f"{nome}: o memorial possui {len(trecho.pontos)} ponto(s) coordenado(s). "
                "Esse trecho não será recalculado nem sobrescrito; se já existir no WebSIG, será preservado."
            )
            continue

        geometria, distancia, duracao, pernas = roteamento_osrm(
            trecho.pontos,
            base_url=osrm_base_url,
            timeout=timeout,
        )
        trecho.geometria = geometria
        trecho.distancia_roteada_km = round(distancia, 3)
        trecho.duracao_roteada_min = round(duracao, 1)
        trecho.distancias_entre_pontos_km = pernas
        alertas_geometria.extend(_avisos_distancias_pontos(nome, trecho))

    km_ida = float(metadados.km_ida) if metadados.km_ida is not None else None
    km_volta = float(metadados.km_volta) if metadados.km_volta is not None else None

    for nome, km_doc in (("IDA", km_ida), ("VOLTA", km_volta)):
        km_roteada = trechos[nome].distancia_roteada_km
        if km_doc and km_roteada:
            diferenca = abs(km_roteada - km_doc) / km_doc
            if diferenca >= 0.15:
                alertas_geometria.append(
                    f"{nome}: a rota calculada ({km_roteada:.2f} km) difere "
                    f"{diferenca * 100:.0f}% da quilometragem do memorial ({km_doc:.2f} km)."
                )

    return MemorialRotaImportado(
        numero_linha=metadados.numero_linha,
        nome_rota=nome_rota_banco(metadados.numero_linha),
        regiao=_extrair_regiao(documento),
        linha=metadados.linha,
        km_ida_documento=km_ida,
        km_volta_documento=km_volta,
        trechos=trechos,
        avisos=avisos,
        alertas_geometria=alertas_geometria,
    )


def salvar_upload_temporario(arquivo) -> Path:
    nome = (arquivo.filename or "").lower()
    if not nome.endswith(".docx"):
        raise ValueError("Envie um memorial no formato DOCX.")
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as temporario:
        arquivo.save(temporario)
        caminho = Path(temporario.name)
    return caminho


def pontos_para_json(pontos: list[dict[str, Any]]) -> str:
    if not isinstance(pontos, list) or len(pontos) < 2:
        raise ValueError("Cada trecho salvo precisa ter pelo menos dois pontos.")
    if len(pontos) > MAX_PONTOS_POR_TRECHO:
        raise ValueError(f"Cada trecho pode ter no máximo {MAX_PONTOS_POR_TRECHO} pontos.")

    campos_texto = ("referencia", "lado_via", "logradouro", "direcao_seguir")
    normalizados: list[dict[str, Any]] = []
    ordens: set[int] = set()

    for ponto in pontos:
        ordem = int(ponto.get("ordem"))
        if ordem in ordens:
            raise ValueError("Há números de ponto repetidos.")
        ordens.add(ordem)

        lat = float(ponto.get("lat"))
        lon = float(ponto.get("lon"))
        if not (-90 <= lat <= 90 and -180 <= lon <= 180):
            raise ValueError("Há coordenadas fora dos limites válidos.")

        normalizado: dict[str, Any] = {
            "ordem": ordem,
            "coordenadas": f"({lat:.6f}, {lon:.6f})",
            "lat": lat,
            "lon": lon,
        }
        for campo in campos_texto:
            normalizado[campo] = texto_limpo(ponto.get(campo))
        if not normalizado["referencia"]:
            raise ValueError(f"O ponto {ordem} não possui referência.")
        normalizados.append(normalizado)

    normalizados.sort(key=lambda item: item["ordem"])
    return json.dumps(normalizados, ensure_ascii=False)


def validar_geojson_linha(
    geometria: Any,
    pontos: list[dict[str, Any]] | None = None,
) -> str:
    if not isinstance(geometria, dict) or geometria.get("type") != "LineString":
        raise ValueError("Geometria de rota inválida.")
    coords = geometria.get("coordinates")
    if not isinstance(coords, list) or len(coords) < 2:
        raise ValueError("A rota precisa ter pelo menos dois vértices.")
    if len(coords) > MAX_VERTICES_GEOMETRIA:
        raise ValueError("A geometria excede o limite de vértices permitido.")

    for coordenada in coords:
        if not isinstance(coordenada, (list, tuple)) or len(coordenada) < 2:
            raise ValueError("Vértice inválido na geometria.")
        lon, lat = float(coordenada[0]), float(coordenada[1])
        if not (-180 <= lon <= 180 and -90 <= lat <= 90):
            raise ValueError("A geometria possui coordenadas fora dos limites válidos.")

    if pontos:
        primeiro = PontoItinerario(1, "", float(pontos[0]["lat"]), float(pontos[0]["lon"]))
        ultimo = PontoItinerario(2, "", float(pontos[-1]["lat"]), float(pontos[-1]["lon"]))
        geom_inicio = PontoItinerario(1, "", float(coords[0][1]), float(coords[0][0]))
        geom_fim = PontoItinerario(2, "", float(coords[-1][1]), float(coords[-1][0]))
        if distancia_haversine_km(primeiro, geom_inicio) > 5.0:
            raise ValueError("O início da geometria está distante do primeiro ponto do memorial.")
        if distancia_haversine_km(ultimo, geom_fim) > 5.0:
            raise ValueError("O fim da geometria está distante do último ponto do memorial.")

    return json.dumps(geometria, ensure_ascii=False)
