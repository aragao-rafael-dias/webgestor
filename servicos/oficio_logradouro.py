from __future__ import annotations

import hashlib
import re
import unicodedata
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from flask import current_app
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    Image,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


MESES = (
    "",
    "janeiro",
    "fevereiro",
    "março",
    "abril",
    "maio",
    "junho",
    "julho",
    "agosto",
    "setembro",
    "outubro",
    "novembro",
    "dezembro",
)

RODAPE_LINHA_1 = (
    "PREFEITURA MUNICIPAL DE LAGARTO, CNPJ: 13.124.052/0001-11"
)
RODAPE_LINHA_2 = (
    "PRAÇA DA PIEDADE, 13, CENTRO, LAGARTO/SE - CEP 49400-000"
)




def _nome_seguro(valor: str) -> str:
    normalizado = unicodedata.normalize(
        "NFKD",
        str(valor or ""),
    ).encode(
        "ascii",
        "ignore",
    ).decode(
        "ascii"
    )
    normalizado = re.sub(
        r"[^A-Za-z0-9_.-]+",
        "_",
        normalizado,
    ).strip("._")
    return normalizado or "documento"

def data_por_extenso(valor: date) -> str:
    return (
        f"Lagarto, {valor.day} de "
        f"{MESES[valor.month]} de {valor.year}."
    )


def _texto(valor: Any, padrao: str = "") -> str:
    resultado = str(valor or "").strip()
    return resultado or padrao


def _cabecalho_path() -> Path:
    configurado = str(
        current_app.config.get(
            "CADASTRO_TERRITORIAL_CABECALHO_OFICIO"
        )
        or ""
    ).strip()

    if configurado:
        caminho = Path(configurado)
    else:
        caminho = (
            Path(current_app.root_path)
            / "static"
            / "img"
            / "cadastro_territorial"
            / "cabecalho_semdu_lagarto.png"
        )

    if not caminho.exists():
        raise FileNotFoundError(
            f"Imagem do cabeçalho não encontrada: {caminho}"
        )

    return caminho


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))

    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)

    shd.set(qn("w:fill"), fill)


def _paragrafo_docx(
    container,
    texto: str = "",
    *,
    negrito: bool = False,
    alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY,
    tamanho: float = 11,
    espaco_depois: float = 6,
    recuo_primeira_linha: float | None = 1.25,
):
    paragrafo = container.add_paragraph()
    paragrafo.alignment = alinhamento
    paragrafo.paragraph_format.space_after = Pt(espaco_depois)
    paragrafo.paragraph_format.line_spacing = 1.15

    if recuo_primeira_linha is not None:
        paragrafo.paragraph_format.first_line_indent = Cm(
            recuo_primeira_linha
        )

    run = paragrafo.add_run(texto)
    run.bold = negrito
    run.font.name = "Arial"
    run.font.size = Pt(tamanho)

    return paragrafo


def _rodape_docx(section) -> None:
    rodape = section.footer
    tabela = rodape.add_table(
        rows=2,
        cols=1,
        width=Cm(18),
    )
    tabela.alignment = 1

    for indice, texto in enumerate(
        (RODAPE_LINHA_1, RODAPE_LINHA_2)
    ):
        paragrafo = tabela.cell(indice, 0).paragraphs[0]
        paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragrafo.paragraph_format.space_after = Pt(0)
        run = paragrafo.add_run(texto)
        run.font.name = "Arial"
        run.font.size = Pt(8)


def _adicionar_tabela_dados_docx(
    documento: Document,
    dados: dict[str, Any],
) -> None:
    linhas = [
        ("Nome proposto", _texto(dados.get("nome_completo"))),
        ("Bairro / localidade", _texto(dados.get("bairro_localidade"), "Não informado")),
        ("Extensão aproximada", f"{float(dados.get('extensao_m') or 0):.2f} m".replace(".", ",")),
        ("Largura prevista", (
            f"{float(dados['largura_m']):.2f} m".replace(".", ",")
            if dados.get("largura_m") is not None
            else "Não informada"
        )),
        ("Coordenada inicial", _texto(dados.get("coordenada_inicial"))),
        ("Coordenada final", _texto(dados.get("coordenada_final"))),
        ("Protocolo", _texto(dados.get("protocolo"))),
    ]

    tabela = documento.add_table(
        rows=1,
        cols=2,
    )
    tabela.style = "Table Grid"
    tabela.autofit = False
    tabela.columns[0].width = Cm(4.5)
    tabela.columns[1].width = Cm(12.5)

    cabecalho = tabela.rows[0].cells
    cabecalho[0].merge(cabecalho[1])
    cabecalho[0].text = "DADOS DO LOGRADOURO PLANEJADO"
    _set_cell_shading(cabecalho[0], "D9E2F3")
    cabecalho[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cabecalho[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for run in cabecalho[0].paragraphs[0].runs:
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9)

    for rotulo, valor in linhas:
        celulas = tabela.add_row().cells
        celulas[0].text = rotulo
        celulas[1].text = valor
        celulas[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        celulas[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        for indice, celula in enumerate(celulas):
            for paragrafo in celula.paragraphs:
                paragrafo.paragraph_format.space_after = Pt(0)
                for run in paragrafo.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
                    run.bold = indice == 0


def gerar_oficio_docx(
    dados: dict[str, Any],
    destino: Path,
) -> None:
    destino.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    documento = Document()
    section = documento.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.1)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.4)
    section.footer_distance = Cm(0.55)

    cabecalho = section.header
    paragrafo_imagem = cabecalho.paragraphs[0]
    paragrafo_imagem.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragrafo_imagem.paragraph_format.space_after = Pt(0)
    paragrafo_imagem.add_run().add_picture(
        str(_cabecalho_path()),
        width=Cm(18.4),
    )

    _rodape_docx(section)

    titulo = documento.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    titulo.paragraph_format.space_after = Pt(8)
    run = titulo.add_run(
        f"Ofício nº {dados['oficio_numero']}/{dados['oficio_ano']} - SEMDU"
    )
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(12)

    data_paragrafo = documento.add_paragraph()
    data_paragrafo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    data_paragrafo.paragraph_format.space_after = Pt(13)
    run = data_paragrafo.add_run(
        data_por_extenso(dados["data_oficio"])
    )
    run.font.name = "Arial"
    run.font.size = Pt(11)

    destinatario = documento.add_paragraph()
    destinatario.alignment = WD_ALIGN_PARAGRAPH.LEFT
    destinatario.paragraph_format.space_after = Pt(12)

    for indice, linha in enumerate(
        [
            "A Sua Excelência o(a) Senhor(a)",
            _texto(dados.get("destinatario_nome"), "RESPONSÁVEL PELA SECRETARIA MUNICIPAL DA FAZENDA"),
            _texto(dados.get("destinatario_cargo"), "Secretário(a) Municipal da Fazenda de Lagarto/SE"),
        ]
    ):
        run = destinatario.add_run(linha)
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.bold = indice > 0
        if indice < 2:
            run.add_break()

    assunto = documento.add_paragraph()
    assunto.alignment = WD_ALIGN_PARAGRAPH.LEFT
    assunto.paragraph_format.space_after = Pt(15)
    run = assunto.add_run("Assunto: ")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run = assunto.add_run(
        _texto(
            dados.get("assunto"),
            "Solicitação de liberação de logradouro planejado.",
        )
    )
    run.font.name = "Arial"
    run.font.size = Pt(11)

    _paragrafo_docx(
        documento,
        _texto(dados.get("vocativo"), "Senhor(a) Secretário(a),"),
        alinhamento=WD_ALIGN_PARAGRAPH.LEFT,
        recuo_primeira_linha=None,
        espaco_depois=12,
    )

    _paragrafo_docx(
        documento,
        (
            "A Secretaria Municipal do Desenvolvimento Urbano - SEMDU, "
            "por meio do módulo Cadastro Territorial do WebSIG Municipal, "
            "encaminha para análise e liberação cadastral o logradouro "
            f"provisoriamente denominado {dados['nome_completo']}, "
            "cujo planejamento técnico foi autorizado por servidor com "
            "função Fiscal da SEMDU."
        ),
    )

    extensao_texto = (
        f"{float(dados.get('extensao_m') or 0):.2f}"
        .replace(".", ",")
    )

    _paragrafo_docx(
        documento,
        (
            "O eixo do logradouro foi definido por levantamento ou importação "
            "geoespacial, encontrando-se registrado no sistema municipal com "
            f"extensão aproximada de {extensao_texto} metros. "
            "Os dados essenciais constam no quadro abaixo."
        ),
    )

    _adicionar_tabela_dados_docx(
        documento,
        dados,
    )

    documento.add_paragraph().paragraph_format.space_after = Pt(0)

    _paragrafo_docx(
        documento,
        "Memorial descritivo:",
        negrito=True,
        alinhamento=WD_ALIGN_PARAGRAPH.LEFT,
        recuo_primeira_linha=None,
        espaco_depois=4,
    )

    _paragrafo_docx(
        documento,
        _texto(dados.get("memorial_descritivo")),
        tamanho=10,
    )

    anexos = dados.get("anexos") or []
    anexos_texto = (
        "; ".join(
            _texto(item.get("nome_original"))
            for item in anexos
            if _texto(item.get("nome_original"))
        )
        or "documentos comprobatórios cadastrados no processo eletrônico"
    )

    _paragrafo_docx(
        documento,
        (
            "Para comprovação da consolidação e instrução da análise, seguem "
            f"vinculados ao processo: {anexos_texto}. A geometria digital e o "
            "histórico de alterações permanecem disponíveis no Cadastro Territorial."
        ),
    )

    _paragrafo_docx(
        documento,
        (
            "Diante do exposto, solicita-se à Secretaria Municipal da Fazenda - "
            "SEMFAZ que proceda à conferência cadastral e, estando os elementos "
            "regulares, altere a situação do logradouro de Planejada - Autorizada "
            "para Planejada - Liberada no sistema compartilhado."
        ),
    )

    _paragrafo_docx(
        documento,
        "Atenciosamente,",
        alinhamento=WD_ALIGN_PARAGRAPH.LEFT,
        recuo_primeira_linha=None,
        espaco_depois=28,
    )

    assinatura = documento.add_paragraph()
    assinatura.alignment = WD_ALIGN_PARAGRAPH.CENTER
    assinatura.paragraph_format.space_after = Pt(0)

    run = assinatura.add_run(
        _texto(dados.get("assinante_nome"), "Responsável pela SEMDU")
    )
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.add_break()

    run = assinatura.add_run(
        _texto(
            dados.get("assinante_cargo"),
            "Secretário(a) Municipal do Desenvolvimento Urbano",
        )
    )
    run.font.name = "Arial"
    run.font.size = Pt(11)

    documento.core_properties.title = (
        f"Ofício {dados['oficio_numero']}/{dados['oficio_ano']} - "
        f"{dados['nome_completo']}"
    )
    documento.core_properties.subject = _texto(dados.get("assunto"))
    documento.core_properties.author = _texto(
        dados.get("assinante_nome"),
        "SEMDU",
    )

    documento.save(destino)


class _DocumentoOficioPDF(BaseDocTemplate):
    def __init__(self, *args, **kwargs):
        self.cabecalho = kwargs.pop("cabecalho")
        super().__init__(*args, **kwargs)

        frame = Frame(
            self.leftMargin,
            self.bottomMargin,
            self.width,
            self.height,
            id="corpo",
            leftPadding=0,
            rightPadding=0,
            topPadding=0,
            bottomPadding=0,
        )

        self.addPageTemplates(
            [
                PageTemplate(
                    id="oficio",
                    frames=[frame],
                    onPage=self._desenhar_pagina,
                )
            ]
        )

    def _desenhar_pagina(self, canvas, doc):
        canvas.saveState()
        largura, altura = A4

        imagem = ImageReader(str(self.cabecalho))
        iw, ih = imagem.getSize()
        max_w = 18.3 * cm
        max_h = 2.45 * cm
        escala = min(max_w / iw, max_h / ih)
        w = iw * escala
        h = ih * escala

        canvas.drawImage(
            imagem,
            (largura - w) / 2,
            altura - 1.0 * cm - h,
            width=w,
            height=h,
            preserveAspectRatio=True,
            mask="auto",
        )

        canvas.setFont("Helvetica", 7.8)
        canvas.setFillColor(colors.HexColor("#111111"))
        canvas.drawCentredString(
            largura / 2,
            1.08 * cm,
            RODAPE_LINHA_1,
        )
        canvas.drawCentredString(
            largura / 2,
            0.72 * cm,
            RODAPE_LINHA_2,
        )

        canvas.restoreState()


def gerar_oficio_pdf(
    dados: dict[str, Any],
    destino: Path,
) -> None:
    destino.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    estilos = getSampleStyleSheet()

    estilo_normal = ParagraphStyle(
        "NormalOficio",
        parent=estilos["BodyText"],
        fontName="Helvetica",
        fontSize=9.7,
        leading=12.3,
        alignment=TA_JUSTIFY,
        firstLineIndent=1.25 * cm,
        spaceAfter=5,
        textColor=colors.HexColor("#111111"),
    )

    estilo_esquerda = ParagraphStyle(
        "EsquerdaOficio",
        parent=estilo_normal,
        alignment=TA_LEFT,
        firstLineIndent=0,
    )

    estilo_direita = ParagraphStyle(
        "DireitaOficio",
        parent=estilo_normal,
        alignment=TA_RIGHT,
        firstLineIndent=0,
    )

    estilo_centro = ParagraphStyle(
        "CentroOficio",
        parent=estilo_normal,
        alignment=TA_CENTER,
        firstLineIndent=0,
    )

    estilo_titulo = ParagraphStyle(
        "TituloOficio",
        parent=estilo_esquerda,
        fontName="Helvetica-Bold",
        fontSize=11.5,
        leading=14,
        spaceAfter=7,
    )

    doc = _DocumentoOficioPDF(
        str(destino),
        pagesize=A4,
        leftMargin=1.8 * cm,
        rightMargin=1.8 * cm,
        topMargin=3.65 * cm,
        bottomMargin=1.55 * cm,
        cabecalho=_cabecalho_path(),
        title=(
            f"Ofício {dados['oficio_numero']}/{dados['oficio_ano']} - "
            f"{dados['nome_completo']}"
        ),
        author=_texto(dados.get("assinante_nome"), "SEMDU"),
    )

    historia = [
        Paragraph(
            f"Ofício nº {dados['oficio_numero']}/{dados['oficio_ano']} - SEMDU",
            estilo_titulo,
        ),
        Paragraph(
            data_por_extenso(dados["data_oficio"]),
            estilo_direita,
        ),
        Spacer(1, 4),
        Paragraph(
            "A Sua Excelência o(a) Senhor(a)<br/>"
            f"<b>{_texto(dados.get('destinatario_nome'), 'RESPONSÁVEL PELA SECRETARIA MUNICIPAL DA FAZENDA')}</b><br/>"
            f"<b>{_texto(dados.get('destinatario_cargo'), 'Secretário(a) Municipal da Fazenda de Lagarto/SE')}</b>",
            estilo_esquerda,
        ),
        Spacer(1, 5),
        Paragraph(
            "<b>Assunto:</b> "
            + _texto(
                dados.get("assunto"),
                "Solicitação de liberação de logradouro planejado.",
            ),
            estilo_esquerda,
        ),
        Spacer(1, 9),
        Paragraph(
            _texto(dados.get("vocativo"), "Senhor(a) Secretário(a),"),
            estilo_esquerda,
        ),
        Paragraph(
            (
                "A Secretaria Municipal do Desenvolvimento Urbano - SEMDU, "
                "por meio do módulo Cadastro Territorial do WebSIG Municipal, "
                "encaminha para análise e liberação cadastral o logradouro "
                f"provisoriamente denominado <b>{dados['nome_completo']}</b>, "
                "cujo planejamento técnico foi autorizado por servidor com "
                "função Fiscal da SEMDU."
            ),
            estilo_normal,
        ),
        Paragraph(
            (
                "O eixo do logradouro foi definido por levantamento ou importação "
                "geoespacial e possui extensão aproximada de "
                f"<b>{float(dados.get('extensao_m') or 0):.2f} metros</b>. "
                "Os dados essenciais constam no quadro abaixo."
            ),
            estilo_normal,
        ),
    ]

    tabela_dados = [
        [
            Paragraph("<b>DADOS DO LOGRADOURO PLANEJADO</b>", estilo_centro),
            "",
        ],
        ["Nome proposto", dados["nome_completo"]],
        ["Bairro / localidade", _texto(dados.get("bairro_localidade"), "Não informado")],
        ["Extensão aproximada", f"{float(dados.get('extensao_m') or 0):.2f} m"],
        [
            "Largura prevista",
            (
                f"{float(dados['largura_m']):.2f} m"
                if dados.get("largura_m") is not None
                else "Não informada"
            ),
        ],
        ["Coordenada inicial", dados["coordenada_inicial"]],
        ["Coordenada final", dados["coordenada_final"]],
        ["Protocolo", dados["protocolo"]],
    ]

    tabela = Table(
        tabela_dados,
        colWidths=[4.4 * cm, 12.0 * cm],
        repeatRows=1,
    )

    tabela.setStyle(
        TableStyle(
            [
                ("SPAN", (0, 0), (1, 0)),
                ("BACKGROUND", (0, 0), (1, 0), colors.HexColor("#D9E2F3")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#7C8495")),
                ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (1, 1), (1, -1), "Helvetica"),
                ("FONTSIZE", (0, 1), (-1, -1), 8.0),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 3.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
            ]
        )
    )

    historia.extend(
        [
            tabela,
            Spacer(1, 9),
            Paragraph("<b>Memorial descritivo:</b>", estilo_esquerda),
            Paragraph(
                _texto(dados.get("memorial_descritivo")),
                estilo_normal,
            ),
        ]
    )

    anexos = dados.get("anexos") or []
    anexos_texto = (
        "; ".join(
            _texto(item.get("nome_original"))
            for item in anexos
            if _texto(item.get("nome_original"))
        )
        or "documentos comprobatórios cadastrados no processo eletrônico"
    )

    historia.extend(
        [
            Paragraph(
                (
                    "Para comprovação da consolidação e instrução da análise, seguem "
                    f"vinculados ao processo: {anexos_texto}. A geometria digital e o "
                    "histórico de alterações permanecem disponíveis no Cadastro Territorial."
                ),
                estilo_normal,
            ),
            Paragraph(
                (
                    "Diante do exposto, solicita-se à Secretaria Municipal da Fazenda - "
                    "SEMFAZ que proceda à conferência cadastral e, estando os elementos "
                    "regulares, altere a situação do logradouro de "
                    "<b>Planejada - Autorizada</b> para "
                    "<b>Planejada - Liberada</b> no sistema compartilhado."
                ),
                estilo_normal,
            ),
            Paragraph("Atenciosamente,", estilo_esquerda),
            Spacer(1, 12),
            Paragraph(
                f"<b>{_texto(dados.get('assinante_nome'), 'Responsável pela SEMDU')}</b><br/>"
                + _texto(
                    dados.get("assinante_cargo"),
                    "Secretário(a) Municipal do Desenvolvimento Urbano",
                ),
                estilo_centro,
            ),
        ]
    )

    doc.build(historia)


def sha256_arquivo(caminho: Path) -> str:
    hash_obj = hashlib.sha256()

    with caminho.open("rb") as arquivo:
        for bloco in iter(
            lambda: arquivo.read(1024 * 1024),
            b"",
        ):
            hash_obj.update(bloco)

    return hash_obj.hexdigest()


def gerar_oficio_logradouro(
    dados: dict[str, Any],
    diretorio: Path,
) -> dict[str, Any]:
    diretorio.mkdir(
        parents=True,
        exist_ok=True,
    )

    base = _nome_seguro(
        f"oficio_{dados['oficio_numero']}_{dados['oficio_ano']}_{dados['protocolo']}"
    )

    caminho_docx = diretorio / f"{base}.docx"
    caminho_pdf = diretorio / f"{base}.pdf"

    gerar_oficio_docx(
        dados,
        caminho_docx,
    )
    gerar_oficio_pdf(
        dados,
        caminho_pdf,
    )

    return {
        "docx": caminho_docx,
        "pdf": caminho_pdf,
        "sha256_docx": sha256_arquivo(caminho_docx),
        "sha256_pdf": sha256_arquivo(caminho_pdf),
    }
